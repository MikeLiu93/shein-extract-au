"""Post-scrape helper: for each row where G=Done and N=empty, search
ebay.com.au via CDP browser (3-tab parallel), extract top-2 Best Match
listings, and write delivered price + cleaned URL to columns N-R.

Usage:
    python ebay_price_check.py                  # picks SUBMITTED_DIR/SHEIN_INPUT_FILENAME
    python ebay_price_check.py "path/to/x.xlsx" # explicit file

Columns written (see docs/superpowers/specs/2026-08-02-ebay-price-check-design.md):
    N: eBay 搜索日期    — YYYY-MM-DD
    O: eBay 同类低价    — numeric (rank-1 or rank-2, whichever is cheaper)
    P: 低价链接         — cleaned eBay itm URL
    Q: eBay 同类高价    — numeric (the pricier of the two)
    R: 高价链接         — cleaned eBay itm URL

Row selection: G exactly 'Done' AND N empty AND K (eBay 标题) non-empty.
"""

import argparse
import logging
import sys
import traceback
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

# Column indices for the eBay result block.
COL_EBAY_SEARCH_DATE = 14
COL_LOW_PRICE        = 15
COL_LOW_URL          = 16
COL_HIGH_PRICE       = 17
COL_HIGH_URL         = 18

logger = logging.getLogger("ebay_price_check")
DEBUG_LOG_DIR = Path(__file__).resolve().parent / "debug_logs"


def _read_ebay_pending_rows(ws) -> list:
    """Return list of dicts for rows that need an eBay search:
    G == 'Done' AND N empty AND K (eBay 标题) non-empty. Non-template
    sheets (missing '链接' in B1) return [].
    Dict shape: {row, seq, ebay_title, shein_title}."""
    # Import here to avoid circular deps at module load time.
    from run_excel import _sheet_matches_template, COL_SEQ, COL_STATUS, COL_SHEIN_TITLE, COL_EBAY_TITLE
    if not _sheet_matches_template(ws):
        return []
    pending = []
    for r in range(2, ws.max_row + 1):
        seq = ws.cell(r, COL_SEQ).value
        status = ws.cell(r, COL_STATUS).value
        search_date = ws.cell(r, COL_EBAY_SEARCH_DATE).value
        ebay_title = ws.cell(r, COL_EBAY_TITLE).value
        shein_title = ws.cell(r, COL_SHEIN_TITLE).value
        if str(status or "") != "Done":
            continue
        if search_date not in (None, ""):
            continue
        if not (ebay_title and str(ebay_title).strip()):
            logger.info("  row %d: skip (eBay 标题 empty)", r)
            continue
        try:
            seq_int = int(seq) if seq is not None else None
        except (TypeError, ValueError):
            seq_int = None
        pending.append({
            "row": r,
            "seq": seq_int,
            "ebay_title": str(ebay_title).strip(),
            "shein_title": str(shein_title or "").strip(),
        })
    return pending


def _write_ebay_result(
    ws, row: int, search_date: str,
    low_price=None, low_url: str = None,
    high_price=None, high_url: str = None,
    no_match: bool = False,
) -> None:
    """Write columns N-R for one row. Always writes N (search_date).
    On no_match=True: O = 'no match', P/Q/R empty. Otherwise: fill the
    slots provided; unspecified stays None."""
    ws.cell(row, COL_EBAY_SEARCH_DATE).value = search_date
    if no_match:
        ws.cell(row, COL_LOW_PRICE).value = "no match"
        return
    if low_price is not None:
        ws.cell(row, COL_LOW_PRICE).value = low_price
    if low_url:
        ws.cell(row, COL_LOW_URL).value = low_url
    if high_price is not None:
        ws.cell(row, COL_HIGH_PRICE).value = high_price
    if high_url:
        ws.cell(row, COL_HIGH_URL).value = high_url


import threading
from concurrent.futures import ThreadPoolExecutor, as_completed

from shein_scraper import (
    CDP_PORT,
    MAX_PARALLEL_TABS,
    RateLimitError,
    _ensure_chrome,
    _inter_url_pause,
    _rate_limit_lock,
    _record_result_for_rate_limit,
    _reset_rate_limit_state,
)
from notify import alert_captcha
from ebay_scraper import search_ebay_au, _ensure_ebay_session
from config import SUBMITTED_DIR, INPUT_FILENAME


def setup_logging():
    DEBUG_LOG_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = DEBUG_LOG_DIR / f"ebay_{ts}.log"
    fmt = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s")

    root = logging.getLogger("ebay_price_check")
    root.handlers.clear()
    root.setLevel(logging.DEBUG)

    fh = logging.FileHandler(log_path, encoding="utf-8")
    fh.setLevel(logging.DEBUG)
    fh.setFormatter(fmt)
    root.addHandler(fh)

    sh = logging.StreamHandler(sys.stdout)
    sh.setLevel(logging.INFO)
    sh.setFormatter(fmt)
    root.addHandler(sh)

    logger.info("Log: %s", log_path)
    return log_path


def safe_save(wb, xlsx_path: Path) -> None:
    """Save workbook. If locked by another user, save as copy with '2' suffix."""
    try:
        wb.save(xlsx_path)
    except PermissionError:
        alt = xlsx_path.with_stem(xlsx_path.stem + "2")
        logger.warning("Cannot save to %s (locked), saving to %s", xlsx_path.name, alt.name)
        wb.save(alt)
        logger.info("Saved to alternate: %s", alt.name)


def _check_one_row(pending_row: dict, port: int) -> dict:
    """Worker: search eBay for one pending row. Returns a result dict with:
      status: 'ok' | 'no_match' | 'captcha' | 'error'
      hits: list[EbayHit] (may be empty)
      row: int (original row for write-back)
      seq: int | None
      error: str | None (populated on captcha/error)
    """
    result = {"row": pending_row["row"], "seq": pending_row.get("seq"),
              "status": "error", "hits": [], "error": None}
    query = pending_row["ebay_title"]
    print(f"[eBay] row {pending_row['row']} seq {pending_row['seq']}: "
          f"searching '{query[:60]}...'")
    try:
        hits = search_ebay_au(query, port)
        if not hits:
            result["status"] = "no_match"
            print("  0 results")
        else:
            result["hits"] = hits
            result["status"] = "ok"
            print(f"  {len(hits)} hit(s): "
                  + ", ".join(f"${h.delivered_price:.2f}" for h in hits))
    except RuntimeError as e:
        msg = str(e)
        if "blocked" in msg.lower() or "captcha" in msg.lower():
            result["status"] = "captcha"
            result["error"] = msg
            print(f"  CAPTCHA: {msg}")
            try:
                alert_captcha(f"eBay AU search: {query[:60]}")
            except Exception:
                pass
        else:
            result["error"] = msg
            print(f"  ERROR: {msg}")
    except Exception as e:
        result["error"] = str(e)
        print(f"  ERROR: {e}")
    return result


def _apply_result_to_row(ws, result: dict, today: str, ws_lock: threading.Lock) -> None:
    """Serialize the write-back (Excel isn't thread-safe)."""
    with ws_lock:
        row = result["row"]
        status = result["status"]
        if status == "ok":
            hits = result["hits"]
            # Sort by delivered_price: cheaper → O/P, pricier → Q/R.
            # Tie-break stable (Best Match rank 1 wins the O slot).
            sorted_hits = sorted(hits, key=lambda h: h.delivered_price)
            low = sorted_hits[0]
            high = sorted_hits[1] if len(sorted_hits) >= 2 else None
            _write_ebay_result(
                ws, row=row, search_date=today,
                low_price=low.delivered_price, low_url=low.url,
                high_price=(high.delivered_price if high else None),
                high_url=(high.url if high else None),
            )
        elif status == "no_match":
            _write_ebay_result(ws, row=row, search_date=today, no_match=True)
        else:
            # captcha / error → do NOT write N; row retries next run.
            logger.info("    row %d skipped (%s)", row, status)


def process_excel(xlsx_path: Path) -> None:
    """Process every worksheet in the file. See module docstring."""
    logger.info("Opening: %s", xlsx_path.name)
    wb = load_workbook(xlsx_path)
    ws_lock = threading.Lock()

    for ws_name in wb.sheetnames:
        ws = wb[ws_name]
        pending = _read_ebay_pending_rows(ws)
        if not pending:
            logger.info("  Sheet '%s': no rows to search (or not template schema)",
                        ws_name.strip())
            continue

        logger.info("  Sheet '%s': %d pending row(s): seq %s",
                    ws_name.strip(), len(pending),
                    [p["seq"] for p in pending])

        _reset_rate_limit_state()
        _ensure_ebay_session(CDP_PORT)

        today = datetime.now().strftime("%Y-%m-%d")
        rate_limited = False

        total = len(pending)
        for batch_start in range(0, total, MAX_PARALLEL_TABS):
            batch_end = min(batch_start + MAX_PARALLEL_TABS, total)
            with _rate_limit_lock:
                if rate_limited:
                    break

            with ThreadPoolExecutor(max_workers=MAX_PARALLEL_TABS) as ex:
                futures = [
                    ex.submit(_check_one_row, pending[i], CDP_PORT)
                    for i in range(batch_start, batch_end)
                ]
                tripped_early = False
                applied = set()  # future ids that have been applied
                for fut in as_completed(futures):
                    try:
                        result = fut.result()
                    except Exception as e:
                        result = {"row": -1, "seq": None, "status": "error",
                                  "hits": [], "error": str(e)}
                    _apply_result_to_row(ws, result, today, ws_lock)
                    applied.add(id(fut))
                    if _record_result_for_rate_limit(
                        "OK" if result["status"] in ("ok", "no_match") else "FAIL"
                    ):
                        tripped_early = True
                        break
                # Drain futures that completed AFTER we broke, but weren't applied yet.
                if tripped_early:
                    for fut in futures:
                        if id(fut) in applied or not fut.done():
                            continue
                        try:
                            result = fut.result()
                            _apply_result_to_row(ws, result, today, ws_lock)
                        except Exception:
                            pass
                    rate_limited = True

            if batch_end < total and not rate_limited:
                _inter_url_pause(batch_end, total)

        safe_save(wb, xlsx_path)
        logger.info("  Saved progress to %s", xlsx_path.name)

    wb.close()
    logger.info("Done: %s", xlsx_path.name)


def main():
    parser = argparse.ArgumentParser(
        description="Post-Shein-scrape eBay price check helper (澳洲站)")
    parser.add_argument("file", nargs="?", default=None,
                        help="Path to .xlsx (default: SHEIN_INPUT_FILENAME under SUBMITTED_DIR)")
    args = parser.parse_args()

    setup_logging()

    if args.file:
        xlsx_path = Path(args.file)
    elif INPUT_FILENAME:
        xlsx_path = SUBMITTED_DIR / INPUT_FILENAME
    else:
        logger.error("No file given and SHEIN_INPUT_FILENAME not set in .env")
        sys.exit(1)

    if not xlsx_path.exists():
        logger.error("File not found: %s", xlsx_path)
        sys.exit(1)

    _ensure_chrome()  # launches or reuses Chrome on CDP_PORT

    try:
        process_excel(xlsx_path)
    except RateLimitError:
        logger.warning("[限流] Rate limited — some rows left unsearched")
    except Exception as e:
        logger.exception("Fatal error: %s", e)
        try:
            tb = traceback.format_exc()
            (DEBUG_LOG_DIR / "last_traceback.txt").write_text(tb, encoding="utf-8")
        except OSError:
            pass
    logger.info("All done.")


if __name__ == "__main__":
    main()
