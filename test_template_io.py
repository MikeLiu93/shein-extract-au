"""Tests for the Chinese-header template reader/writer in run_excel.py."""
from openpyxl import Workbook

from run_excel import _read_pending_rows, _write_result_row


CN_HEADERS = ["编号", "链接", "原价", "运费", "变体",
              "日期", "状态", "希音价格", "希音标题",
              "eBay标题", "eBay价格", "库存"]


def _blank_ws(rows: list):
    wb = Workbook()
    ws = wb.active
    ws.title = "ZR1"
    for ci, h in enumerate(CN_HEADERS, 1):
        ws.cell(1, ci).value = h
    for ri, r in enumerate(rows, 2):
        for ci, v in enumerate(r, 1):
            ws.cell(ri, ci).value = v
    return ws


# ── _read_pending_rows ──────────────────────────────────────────────────────

def test_reads_only_rows_ready_and_not_done():
    # Row 2: ready. Row 3: has date → skip. Row 4: no URL → skip.
    ws = _blank_ws([
        [1, "http://u1", 10.0, 7.95, "",  None,        None, None, None, None, None, None],
        [2, "http://u2", 12.0, 7.95, "",  "2026-07-30", "Done", None, None, None, None, None],
        [3, None,        15.0, 7.95, "",  None,        None, None, None, None, None, None],
    ])
    pending = _read_pending_rows(ws)
    assert len(pending) == 1
    p = pending[0]
    assert p["row"] == 2
    assert p["seq"] == 1
    assert p["url"] == "http://u1"
    assert p["price"] == 10.0
    assert p["shipping"] == 7.95
    assert p["variant_filter"] == ""


def test_blank_price_and_shipping_kept_as_none():
    ws = _blank_ws([
        [1, "http://u1", None, None, "", None, None, None, None, None, None, None],
    ])
    pending = _read_pending_rows(ws)
    assert len(pending) == 1
    p = pending[0]
    assert p["price"] is None
    assert p["shipping"] is None


def test_variant_filter_captured():
    ws = _blank_ws([
        [1, "http://u1", 10.0, 7.95, "Black, Red / M, L", None, None, None, None, None, None, None],
    ])
    pending = _read_pending_rows(ws)
    assert pending[0]["variant_filter"] == "Black, Red / M, L"


def test_rejects_non_template_sheet():
    # A sheet missing the Chinese '链接' header in col B is not our schema.
    wb = Workbook()
    ws = wb.active
    ws.cell(1, 1).value = "Seq"
    ws.cell(1, 2).value = "Website"  # English → reject
    ws.cell(2, 1).value = 1
    ws.cell(2, 2).value = "http://u"
    pending = _read_pending_rows(ws)
    assert pending == []  # empty list = "not our schema, skip"


# ── _write_result_row ───────────────────────────────────────────────────────

def test_write_populates_result_columns():
    ws = _blank_ws([
        [1, "http://u1", 10.0, 7.95, "", None, None, None, None, None, None, None],
    ])
    _write_result_row(
        ws, row=2,
        date="2026-08-01",
        status="Done",
        web_price="$9.99–$14.99",
        shein_title="Foo Product",
        ebay_title="NEW Foo Product Adjustable Waterproof",
        ebay_price=27.95,
        stock="Black-M: 12 / Black-L: 缺货",
    )
    assert ws.cell(2, 6).value == "2026-08-01"
    assert ws.cell(2, 7).value == "Done"
    assert ws.cell(2, 8).value == "$9.99–$14.99"
    assert ws.cell(2, 9).value == "Foo Product"
    assert ws.cell(2, 10).value == "NEW Foo Product Adjustable Waterproof"
    assert ws.cell(2, 11).value == 27.95
    assert ws.cell(2, 12).value == "Black-M: 12 / Black-L: 缺货"


def test_write_only_date_status_on_failure():
    ws = _blank_ws([
        [1, "http://u1", 10.0, 7.95, "", None, None, None, None, None, None, None],
    ])
    _write_result_row(ws, row=2, date="2026-08-01", status="Failed")
    assert ws.cell(2, 6).value == "2026-08-01"
    assert ws.cell(2, 7).value == "Failed"
    # Result cols untouched
    assert ws.cell(2, 8).value is None
    assert ws.cell(2, 9).value is None
    assert ws.cell(2, 10).value is None
    assert ws.cell(2, 11).value is None
    assert ws.cell(2, 12).value is None


if __name__ == "__main__":
    test_reads_only_rows_ready_and_not_done()
    test_blank_price_and_shipping_kept_as_none()
    test_variant_filter_captured()
    test_rejects_non_template_sheet()
    test_write_populates_result_columns()
    test_write_only_date_status_on_failure()
    print("ALL PASS")
