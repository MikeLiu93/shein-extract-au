"""Tests for ebay_price_check row reader (_read_ebay_pending_rows) and
writer (_write_ebay_result) against the template schema."""
from openpyxl import Workbook

from ebay_price_check import _read_ebay_pending_rows, _write_ebay_result


CN_HEADERS = ["编号", "链接", "原价", "运费", "变体",
              "日期", "状态", "图片", "希音价格", "希音标题",
              "eBay标题", "eBay价格", "库存",
              "eBay搜索日期", "eBay同类低价", "低价链接",
              "eBay同类高价", "高价链接",
              "Shein重跑日期", "更新价格", "更新库存"]
NCOLS = len(CN_HEADERS)  # 21


def _make_ws(rows: list):
    wb = Workbook()
    ws = wb.active
    ws.title = "ZR1"
    for ci, h in enumerate(CN_HEADERS, 1):
        ws.cell(1, ci).value = h
    for ri, r in enumerate(rows, 2):
        padded = list(r) + [None] * (NCOLS - len(r))
        for ci, v in enumerate(padded, 1):
            ws.cell(ri, ci).value = v
    return ws


# ── _read_ebay_pending_rows ─────────────────────────────────────────────────

def test_reads_done_rows_with_empty_search_date():
    # Row 2: Done + no search date → PICK.
    # Row 3: Failed → skip.
    # Row 4: Done + already searched → skip.
    # Row 5: Delisted → skip.
    ws = _make_ws([
        [1, "url1", 10, 7.95, "", "2026-08-02", "Done",
         None, 9.99, "shein-title", "ebay-title", 27.95, "M: 20", None],
        [2, "url2", 12, 0,    "", "2026-08-02", "Failed",
         None, None, None, None, None, None, None],
        [3, "url3", 15, 7.95, "", "2026-08-02", "Done",
         None, 14.99, "st3", "et3", 37.93, "one-size: 20", "2026-08-02"],  # N filled
        [4, "url4", 20, 7.95, "", "2026-08-02", "Delisted",
         None, None, None, None, None, None, None],
    ])
    pending = _read_ebay_pending_rows(ws)
    assert len(pending) == 1
    p = pending[0]
    assert p["row"] == 2
    assert p["seq"] == 1
    assert p["ebay_title"] == "ebay-title"
    assert p["shein_title"] == "shein-title"


def test_skips_row_with_empty_ebay_title():
    # G=Done, N=empty, but K (eBay 标题) is empty → skip (can't search).
    ws = _make_ws([
        [1, "url1", 10, 7.95, "", "2026-08-02", "Done",
         None, 9.99, "shein-title", None, 27.95, "M: 20", None],
    ])
    pending = _read_ebay_pending_rows(ws)
    assert pending == []


def test_case_sensitive_done_only():
    # 'done' (lowercase) → skip; only exact 'Done' picks.
    ws = _make_ws([
        [1, "url1", 10, 7.95, "", "2026-08-02", "done",
         None, 9.99, "st", "et", 27.95, "M: 20", None],
        [2, "url2", 10, 7.95, "", "2026-08-02", "Done",
         None, 9.99, "st", "et", 27.95, "M: 20", None],
    ])
    pending = _read_ebay_pending_rows(ws)
    assert [p["seq"] for p in pending] == [2]


def test_rejects_non_template_sheet():
    wb = Workbook()
    ws = wb.active
    ws.cell(1, 1).value = "Seq"
    ws.cell(1, 2).value = "Website"  # English → not our schema
    pending = _read_ebay_pending_rows(ws)
    assert pending == []


# ── _write_ebay_result ──────────────────────────────────────────────────────

def test_write_two_results():
    ws = _make_ws([
        [1, "url1", 10, 7.95, "", "2026-08-02", "Done",
         None, 9.99, "st", "et", 27.95, "M: 20", None],
    ])
    _write_ebay_result(
        ws, row=2,
        search_date="2026-08-02",
        low_price=25.90, low_url="https://www.ebay.com.au/itm/111",
        high_price=32.50, high_url="https://www.ebay.com.au/itm/222",
    )
    assert ws.cell(2, 14).value == "2026-08-02"           # N
    assert ws.cell(2, 15).value == 25.90                   # O
    assert ws.cell(2, 16).value == "https://www.ebay.com.au/itm/111"  # P
    assert ws.cell(2, 17).value == 32.50                   # Q
    assert ws.cell(2, 18).value == "https://www.ebay.com.au/itm/222"  # R


def test_write_one_result_leaves_high_empty():
    ws = _make_ws([[1, "url1", 10, 7.95, "", "2026-08-02", "Done",
                    None, 9.99, "st", "et", 27.95, "M: 20", None]])
    _write_ebay_result(ws, row=2, search_date="2026-08-02",
                       low_price=25.90, low_url="https://www.ebay.com.au/itm/111")
    assert ws.cell(2, 14).value == "2026-08-02"
    assert ws.cell(2, 15).value == 25.90
    assert ws.cell(2, 16).value == "https://www.ebay.com.au/itm/111"
    assert ws.cell(2, 17).value is None
    assert ws.cell(2, 18).value is None


def test_write_zero_result_marks_no_match():
    ws = _make_ws([[1, "url1", 10, 7.95, "", "2026-08-02", "Done",
                    None, 9.99, "st", "et", 27.95, "M: 20", None]])
    _write_ebay_result(ws, row=2, search_date="2026-08-02", no_match=True)
    assert ws.cell(2, 14).value == "2026-08-02"
    assert ws.cell(2, 15).value == "no match"
    assert ws.cell(2, 16).value is None
    assert ws.cell(2, 17).value is None
    assert ws.cell(2, 18).value is None


if __name__ == "__main__":
    test_reads_done_rows_with_empty_search_date()
    test_skips_row_with_empty_ebay_title()
    test_case_sensitive_done_only()
    test_rejects_non_template_sheet()
    test_write_two_results()
    test_write_one_result_leaves_high_empty()
    test_write_zero_result_marks_no_match()
    print("ALL PASS")
