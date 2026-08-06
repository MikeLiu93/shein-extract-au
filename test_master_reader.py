"""Tests for _read_master_pending_rows — reads the 6-col master schema and
returns rows where F (是否要跑) normalizes to 'Y' (case + whitespace tolerant)."""
from openpyxl import Workbook

from run_excel import _read_master_pending_rows


MASTER_HEADERS = ["编号", "链接", "原价", "运费", "变体", "是否要跑"]
NCOLS = len(MASTER_HEADERS)  # 6


def _make_ws(rows: list):
    wb = Workbook()
    ws = wb.active
    ws.title = "ZR1"
    for ci, h in enumerate(MASTER_HEADERS, 1):
        ws.cell(1, ci).value = h
    for ri, r in enumerate(rows, 2):
        padded = list(r) + [None] * (NCOLS - len(r))
        for ci, v in enumerate(padded, 1):
            ws.cell(ri, ci).value = v
    return ws


def test_y_uppercase_triggers():
    ws = _make_ws([
        [1, "http://u1", 10.0, 7.95, "", "Y"],
    ])
    pending = _read_master_pending_rows(ws)
    assert len(pending) == 1
    assert pending[0]["row"] == 2
    assert pending[0]["seq"] == 1
    assert pending[0]["url"] == "http://u1"
    assert pending[0]["price"] == 10.0
    assert pending[0]["shipping"] == 7.95
    assert pending[0]["variant_filter"] == ""


def test_y_lowercase_and_whitespace_triggers():
    ws = _make_ws([
        [1, "http://u1", 10.0, 7.95, "", "y"],
        [2, "http://u2", 10.0, 7.95, "", "  Y  "],
        [3, "http://u3", 10.0, 7.95, "", "yes"],  # only exact Y (any case) triggers
    ])
    pending = _read_master_pending_rows(ws)
    seqs = [p["seq"] for p in pending]
    assert seqs == [1, 2], seqs  # 'yes' does NOT trigger; only 'y' / 'Y'


def test_non_y_is_skipped():
    ws = _make_ws([
        [1, "http://u1", 10.0, 7.95, "", ""],       # empty
        [2, "http://u2", 10.0, 7.95, "", "N"],       # explicit skip
        [3, "http://u3", 10.0, 7.95, "", "完成"],    # arbitrary marker
        [4, "http://u4", 10.0, 7.95, "", None],      # None
    ])
    pending = _read_master_pending_rows(ws)
    assert pending == []


def test_variant_filter_and_optional_price_shipping():
    ws = _make_ws([
        [1, "http://u1", None, None, "Black, Red / M, L", "Y"],
    ])
    pending = _read_master_pending_rows(ws)
    assert len(pending) == 1
    p = pending[0]
    assert p["price"] is None
    assert p["shipping"] is None
    assert p["variant_filter"] == "Black, Red / M, L"


def test_non_master_sheet_returns_empty():
    # A sheet whose B1 is not '链接' isn't the master schema.
    wb = Workbook()
    ws = wb.active
    ws.cell(1, 1).value = "Seq"
    ws.cell(1, 2).value = "Website"
    ws.cell(2, 1).value = 1
    ws.cell(2, 6).value = "Y"
    assert _read_master_pending_rows(ws) == []


def test_non_numeric_seq_skipped():
    ws = _make_ws([
        ["oops", "http://u1", 10.0, 7.95, "", "Y"],
        [2, "http://u2", 10.0, 7.95, "", "Y"],
    ])
    pending = _read_master_pending_rows(ws)
    assert [p["seq"] for p in pending] == [2]


def test_missing_url_skipped():
    ws = _make_ws([
        [1, None, 10.0, 7.95, "", "Y"],
        [2, "http://u2", 10.0, 7.95, "", "Y"],
    ])
    pending = _read_master_pending_rows(ws)
    assert [p["seq"] for p in pending] == [2]


if __name__ == "__main__":
    test_y_uppercase_triggers()
    test_y_lowercase_and_whitespace_triggers()
    test_non_y_is_skipped()
    test_variant_filter_and_optional_price_shipping()
    test_non_master_sheet_returns_empty()
    test_non_numeric_seq_skipped()
    test_missing_url_skipped()
    print("ALL PASS")
